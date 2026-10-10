from types import SimpleNamespace
import pytest
from aquillm.ingestion.figure_extraction import office, spreadsheet, office_convert
from aquillm.ingestion.figure_extraction.types import ExtractedFigure


def figure(data):
    return ExtractedFigure(image_bytes=data, image_format='png', figure_index=0, nearby_text='', width=100, height=100, location_metadata={})


@pytest.mark.parametrize('format', ['docx', 'pptx', 'xlsx', 'ods'])
@pytest.mark.parametrize('status', ['unavailable', 'failed', 'converted_empty', 'converted_figures', 'direct_failure'])
def test_conversion_outcomes_always_attempt_direct_and_deduplicate(monkeypatch, format, status):
    module = office if format in ('docx', 'pptx') else spreadsheet
    conversion = SimpleNamespace(status='converted' if status.startswith('converted') or status == 'direct_failure' else status, pdf_bytes=b'pdf')
    monkeypatch.setattr(office_convert, 'convert_figures_pdf', lambda *a: conversion, raising=False)
    import aquillm.ingestion.figure_extraction.pdf as pdf
    monkeypatch.setattr(pdf, 'extract_figures', lambda *a: iter([figure(b'shared')]) if status in ('converted_figures', 'direct_failure') else iter([]))
    def direct(*args):
        if status == 'direct_failure':
            raise ValueError('synthetic direct failure')
        yield figure(b'shared')
        yield figure(b'direct')
    monkeypatch.setattr(module, '_extract_' + format + '_direct', direct)
    output = list(getattr(module, 'extract_figures_' + format)(b'office', 'test.' + format))
    assert [f.image_bytes for f in output] == ([b'shared'] if status == 'direct_failure' else [b'shared', b'direct'])


@pytest.mark.parametrize('available,result,status', [(False,None,'unavailable'), (True,None,'failed'), (True,b'pdf','converted')])
def test_conversion_is_a_non_generator_explicit_outcome(monkeypatch, available, result, status):
    monkeypatch.setattr(office_convert, 'is_libreoffice_available', lambda: available)
    monkeypatch.setattr(office_convert, 'convert_to_pdf', lambda *a: result)
    assert hasattr(office_convert, 'convert_figures_pdf'), 'explicit conversion outcome missing'
    output = office_convert.convert_figures_pdf(b'office', 'docx', 'test.docx')
    assert output.status == status
    assert output.pdf_bytes == (result if status == 'converted' else None)


def image_bytes():
    import io
    import random
    from PIL import Image
    output = io.BytesIO()
    image = Image.frombytes('RGB', (128, 128), random.Random(42).randbytes(128 * 128 * 3))
    image.save(output, format='PNG')
    image.close()
    return output.getvalue()


@pytest.mark.parametrize('format', ['docx', 'pptx'])
def test_real_embedded_images_survive_unavailable_conversion_and_grouping(monkeypatch, format):
    import io
    png = image_bytes()
    stream = io.BytesIO()
    monkeypatch.setattr(office_convert, 'convert_figures_pdf', lambda *a: SimpleNamespace(status='unavailable', pdf_bytes=None))
    if format == 'docx':
        from docx import Document
        document = Document()
        document.add_paragraph('Native fact')
        document.add_picture(io.BytesIO(png))
    else:
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(io.BytesIO(png), 0, 0)
    document.save(stream)
    output = list(getattr(office, 'extract_figures_' + format)(stream.getvalue(), 'test.' + format))
    assert [f.image_bytes for f in output] == [png]


def test_direct_extractor_failure_at_call_still_preserves_converted_figures(monkeypatch):
    monkeypatch.setattr(office_convert, 'convert_figures_pdf', lambda *a: SimpleNamespace(status='converted', pdf_bytes=b'pdf'))
    import aquillm.ingestion.figure_extraction.pdf as pdf
    monkeypatch.setattr(pdf, 'extract_figures', lambda *a: iter([figure(b'converted')]))
    def fail(*args):
        raise ValueError('synthetic call failure')
    monkeypatch.setattr(office, '_extract_docx_direct', fail)
    assert [f.image_bytes for f in office.extract_figures_docx(b'office')] == [b'converted']
