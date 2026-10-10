import importlib
from dataclasses import replace
import json
from hashlib import sha256

import pytest
from lib.parsers.source_contracts import SourceContractError
from .test_source_artifacts import artifact_storage


def modules():
    return (importlib.import_module('apps.ingestion.services.extraction_revisions'), importlib.import_module('lib.parsers.canonical_projection'))


@pytest.mark.django_db
def test_payload_binder_persists_exact_immutable_revision_and_retries(artifact_storage):
    from aquillm.ingestion.parsers import extract_text_payloads
    from apps.ingestion.models import ExtractionRevision
    service, projection = modules()
    payload = extract_text_payloads('data.csv', b'Product,Domestic,Foreign\nWidget,,7\n')[0]
    revision_id = service.bind_payload_extraction(payload, using='default')
    revision = ExtractionRevision.objects.get(pk=revision_id)
    assert payload.source_occurrence.occurrence_id == revision.occurrence_id
    assert payload.extraction_revision_id == revision_id
    assert revision.canonical_text == payload.full_text == payload.extraction.canonical_text
    assert revision.canonical_sha256 == sha256(payload.full_text.encode()).hexdigest()
    assert projection.elements_from_json(revision.ordered_elements) == payload.extraction.elements
    manifest = projection.extraction_manifest(payload.extraction)
    assert revision.manifest_sha256 == sha256(json.dumps(manifest, sort_keys=True, separators=(',', ':'), ensure_ascii=True, allow_nan=False).encode()).hexdigest()
    assert service.bind_payload_extraction(payload, using='default') == revision_id
    assert ExtractionRevision.objects.count() == 1


@pytest.mark.django_db
def test_revision_rejects_forged_occurrence_and_artifact(artifact_storage):
    from aquillm.ingestion.parsers import extract_text_payloads
    service, _ = modules()
    payload = extract_text_payloads('data.csv', b'A,B\nx,7\n')[0]
    service.bind_payload_extraction(payload, using='default')
    occurrence = payload.source_occurrence
    with pytest.raises(SourceContractError):
        service.persist_extraction_revision(replace(occurrence, occurrence_key='0' * 64), payload.extraction, using='default')
    with pytest.raises(SourceContractError):
        service.persist_extraction_revision(occurrence, replace(payload.extraction, artifact=replace(occurrence.artifact, sha256='0' * 64)), using='default')


@pytest.mark.django_db
def test_binder_rejects_bool_as_existing_revision_id(artifact_storage):
    from aquillm.ingestion.parsers import extract_text_payloads
    service, _ = modules()
    payload = extract_text_payloads('data.csv', b'A,B\nx,7\n')[0]
    service.bind_payload_extraction(payload, using='default')
    payload.extraction_revision_id = True
    with pytest.raises(SourceContractError, match='extraction_revision_binding_mismatch'):
        service.bind_payload_extraction(payload, using='default')


def test_serializer_rejects_unknown_fields_bool_offsets_and_mutable_cells():
    from lib.parsers.source_contracts import CellV1, SourceElementV1
    _, projection = modules()
    element = SourceElementV1('row', 'table', 0, 1, None, 'Sheet', None, None, (CellV1(1, 1, 'A1', 'x', None, 'literal'),))
    wire = projection.elements_to_json((element,))
    assert set(wire[0]) == {'element_key','kind','start','end','page','sheet','slide','parent_key','cells','bbox'}
    assert wire[0]['page'] is None
    assert projection.elements_from_json(wire) == (element,)
    for changed in [dict(wire[0], unknown=True), dict(wire[0], start=False)]:
        with pytest.raises(SourceContractError):
            projection.elements_from_json([changed])
    with pytest.raises(SourceContractError):
        projection.elements_to_json((replace(element, cells=list(element.cells)),))
    with pytest.raises(SourceContractError):
        projection.elements_from_json([dict(wire[0], kind=['table'])])


@pytest.mark.django_db
def test_identical_archive_members_keep_distinct_revision_occurrences(artifact_storage):
    import io
    import zipfile
    from aquillm.ingestion.parsers import extract_text_payloads
    from apps.ingestion.models import ExtractionRevision
    service, _ = modules()
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w') as archive:
        archive.writestr('data.csv', b'A,B\nx,7\n')
        with pytest.warns(UserWarning):
            archive.writestr('data.csv', b'A,B\nx,7\n')
    outputs = extract_text_payloads('archive.zip', stream.getvalue())
    ids = [service.bind_payload_extraction(payload, using='default') for payload in outputs]
    assert ids[0] != ids[1]
    assert ExtractionRevision.objects.get(pk=ids[0]).canonical_sha256 == ExtractionRevision.objects.get(pk=ids[1]).canonical_sha256
    assert outputs[0].source_occurrence.occurrence_key != outputs[1].source_occurrence.occurrence_key


@pytest.mark.django_db
@pytest.mark.parametrize('field,bad_value', [('parser_key', 'wrong-parser'), ('manifest_sha256', '0' * 64), ('occurrence_id', 999999), ('canonical_text', 'corrupted snapshot'), ('canonical_sha256', '0' * 64), ('ordered_elements', []), ('quality', 'unsupported'), ('issue_codes', ['incorrect'])])
def test_revision_collision_verifies_every_persisted_field(artifact_storage, field, bad_value):
    from unittest.mock import patch
    from apps.ingestion.models import ExtractionRevision
    from aquillm.ingestion.parsers import extract_text_payloads
    service, _ = modules()
    payload = extract_text_payloads('data.csv', b'A,B\nx,7\n')[0]
    revision_id = service.bind_payload_extraction(payload, using='default')
    # The real DB trigger rejects raw UPDATE, so simulate a corrupted ORM
    # hydration without weakening that trigger or mutating the persisted row.
    hydrated = ExtractionRevision.objects.get(pk=revision_id)
    setattr(hydrated, field, bad_value)
    with patch.object(ExtractionRevision.objects, 'using') as manager:
        manager.return_value.get_or_create.return_value = (hydrated, False)
        with pytest.raises(SourceContractError, match='extraction_revision_registration_mismatch'):
            service.persist_extraction_revision(payload.source_occurrence, payload.extraction, using='default')


@pytest.mark.django_db
def test_generic_text_projection_is_explicitly_partial_without_invented_structure(artifact_storage):
    from aquillm.ingestion.parsers import extract_text_payloads
    service, _ = modules()
    payload = extract_text_payloads('data.json', '{"metric": "−4 kg"}'.encode())[0]
    assert payload.extraction.quality == 'partial'
    assert payload.extraction.issue_codes == ('structure_unavailable',)
    assert payload.extraction.canonical_text == payload.full_text
    assert payload.extraction.elements[0].cells == ()
    assert service.bind_payload_extraction(payload, using='default') > 0


@pytest.mark.django_db
@pytest.mark.parametrize('format', ['docx', 'xlsx'])
def test_parser_policy_v2_preserves_frozen_v1_identity_and_replays_both(artifact_storage, format):
    import io
    from aquillm.ingestion.parsers import extract_text_payloads
    from apps.ingestion.models import ExtractionRevision
    from apps.ingestion.services.source_artifacts import bind_payload_source
    stream = io.BytesIO()
    if format == 'docx':
        from docx import Document
        from lib.parsers.tests.test_structural_extraction import _wrap_docx_control
        doc = Document()
        doc.add_paragraph('Before')
        _wrap_docx_control(doc.add_paragraph('ControlledRevenue 93017')._p)
        doc.add_paragraph('After')
        doc.save(stream)
    else:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        book = Workbook()
        book.active.append(['Metric', 'Value'])
        book.active.append(['A', 7])
        chart = BarChart()
        chart.title = 'ChartOnlyRevenue 93017'
        chart.add_data(Reference(book.active, min_col=2, min_row=1, max_row=2), titles_from_data=True)
        book.active.add_chart(chart, 'D1')
        book.save(stream)
        book.close()
    service, _ = modules()
    payload = extract_text_payloads('policy.' + format, stream.getvalue())[0]
    current = payload.extraction
    assert current.parser_key == f'native-{format}-ordered-v2'
    # These capture the original v1 behavior for the exact same source bytes:
    # body controls were omitted, and a cell-only chart extraction was complete.
    if format == 'docx':
        legacy = replace(current, parser_key='native-docx-ordered-v1', canonical_text='Before\nAfter',
                         elements=(current.elements[0], replace(current.elements[-1], element_key='element-2', start=7, end=12)))
    else:
        legacy = replace(current, parser_key='native-xlsx-ordered-v1', quality='complete', issue_codes=())
    occurrence = bind_payload_source(payload, using='default')
    old_id = service.persist_extraction_revision(occurrence, legacy, using='default')
    new_id = service.bind_payload_extraction(payload, using='default')
    assert new_id != old_id
    assert service.persist_extraction_revision(occurrence, legacy, using='default') == old_id
    assert service.persist_extraction_revision(occurrence, current, using='default') == new_id
    old, new = ExtractionRevision.objects.get(pk=old_id), ExtractionRevision.objects.get(pk=new_id)
    assert old.parser_key == f'native-{format}-ordered-v1'
    assert new.parser_key == current.parser_key
    assert old.canonical_text == legacy.canonical_text
    assert old.quality == legacy.quality
    assert old.issue_codes == list(legacy.issue_codes)
    assert old.manifest_sha256 != new.manifest_sha256
    assert old.occurrence_id == new.occurrence_id == occurrence.occurrence_id
    assert ExtractionRevision.objects.count() == 2


@pytest.mark.django_db
def test_figure_payload_uses_parent_extraction_and_its_own_exact_media(artifact_storage, monkeypatch):
    from aquillm.ingestion.parsers import extract_text_payloads
    from aquillm.ingestion.types import ExtractedTextPayload
    from .test_archive_source_binding import pdf_bytes
    from apps.ingestion.models import ExtractionRevision
    service, _ = modules()
    def figures(filename, data, format, outputs):
        outputs.append(ExtractedTextPayload('figure', 'document_figure', 'synthetic caption', modality='image', media_bytes=b'picture'))
    monkeypatch.setattr('aquillm.ingestion.parsers.extract_figure_payloads_for_format', figures)
    body, figure = extract_text_payloads('paper.pdf', pdf_bytes(150))
    assert figure.extraction is body.extraction
    body_id = service.bind_payload_extraction(body, using='default')
    figure_id = service.bind_payload_extraction(figure, using='default')
    assert body_id == figure_id
    assert ExtractionRevision.objects.get(pk=figure_id).canonical_text == body.full_text
    assert figure.media_artifact.sha256 == sha256(b'picture').hexdigest()
    assert figure.media_artifact != figure.source_artifact
