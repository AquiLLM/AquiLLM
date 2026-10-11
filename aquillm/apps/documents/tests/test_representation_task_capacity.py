"""Actual native DOCX input and durable task failure, without a partial publication."""
import io
from datetime import timedelta
from uuid import uuid4

import pytest
from django.apps import apps
from django.utils import timezone
from docx import Document

from apps.documents.models import ChunkPublication, DocumentRepresentationState
from apps.documents.tasks.chunking import create_chunks
from aquillm.task_ingest_uploaded import run_ingest_uploaded_file
from .representation_support import upload as upload
from .test_representation_tasks import delivery

pytestmark = pytest.mark.django_db(transaction=True)


def test_native_docx_work_limit_keeps_exact_original_revision_and_prior_active(upload, monkeypatch):
    word = Document()
    word.add_heading("Heading of 20 chars!", level=1)
    word.add_paragraph("x" * 1500)
    stream = io.BytesIO()
    word.save(stream)
    item = upload((("native.docx", stream.getvalue()),))
    run_ingest_uploaded_file(item.pk)
    output = item.outputs.get(status="success", output_kind="text")
    doc = apps.get_model(output.document_model).objects.get(pk=output.document_pkid)
    assert create_chunks.run(**delivery(doc)) == "activated"
    doc.refresh_from_db()
    state = DocumentRepresentationState.objects.get(document_id=doc.id)
    original_key = state.active_generation_id
    original_ids = list(doc.chunks.values_list("pk", flat=True))
    config = apps.get_app_config("aquillm")
    monkeypatch.setattr(config, "chunk_size", 20)
    monkeypatch.setattr(config, "chunk_overlap", 0)
    intent = ChunkPublication.objects.create(concrete_model_label=doc._meta.label_lower,
        document_pkid=doc.pkid, document_id=doc.id, source_hash=doc.full_text_hash,
        collection_ids=[doc.collection_id], generation=uuid4(), attempts=1,
        next_attempt_at=timezone.now() + timedelta(seconds=900))
    assert create_chunks.run(**delivery(doc)) == "failed"
    intent.refresh_from_db()
    state.refresh_from_db()
    failed = state.generations.get(status="failed")
    assert failed.failure_code == "chunker_work_limit" and failed.expected_count == 0
    assert failed.manifest["chunks"] == [] and failed.extraction_revision_id == output.extraction_revision_id
    assert failed.parser_key == "native-docx-ordered-v2"
    assert intent.failure_kind == "contract" and intent.last_error == "chunker_work_limit"
    assert state.active_generation_id == original_key
    assert list(doc.chunks.values_list("pk", flat=True)) == original_ids
