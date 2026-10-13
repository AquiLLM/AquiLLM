"""Bounded, deterministic public source bytes, independent of expected answers."""

import io
import re
import warnings
from zipfile import ZIP_STORED, ZipFile, ZipInfo

MAX_SOURCE_BYTES = 1_000_000


def _zip(entries):
    output = io.BytesIO()
    with ZipFile(output, 'w', compression=ZIP_STORED) as archive:
        for filename, data in entries:
            info = ZipInfo(filename, date_time=(2020, 1, 1, 0, 0, 0))
            # Duplicate exact names are intentional occurrence fixtures.
            with warnings.catch_warnings():
                warnings.filterwarnings('ignore', message='Duplicate name:', category=UserWarning)
                archive.writestr(info, data)
    return output.getvalue()


def _stable_office(data):
    with ZipFile(io.BytesIO(data)) as archive:
        entries = []
        for info in archive.infolist():
            body = archive.read(info)
            if info.filename == 'docProps/core.xml':
                body = re.sub(rb'\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ',
                              b'2020-01-01T00:00:00Z', body)
            entries.append((info.filename, body))
    return _zip(entries)


def build_source(spec):
    """Construct source only from source fields, never fixture expectations."""
    repeat = spec.get('repeat', 0)
    if type(repeat) is not int or not 0 <= repeat <= 1000:
        raise ValueError('fixture_source_limit')
    if spec['format'] == 'text':
        data = (spec.get('prefix', '') * repeat + spec['text']).encode('utf-8')
    elif spec['format'] == 'zip':
        if len(spec['members']) > 32:
            raise ValueError('fixture_source_limit')
        data = _zip((member['filename'], member['text'].encode('utf-8'))
                    for member in spec['members'])
    elif spec['format'] in ('xlsx', 'docx'):
        rows = spec['rows']
        if len(rows) > 100 or any(len(row) > 32 for row in rows):
            raise ValueError('fixture_source_limit')
        stream = io.BytesIO()
        if spec['format'] == 'xlsx':
            from openpyxl import Workbook
            document = Workbook()
            document.active.title = spec['sheet']
            for row in rows:
                document.active.append(row)
            document.save(stream)
            document.close()
        else:
            from docx import Document
            document = Document()
            document.add_heading(spec['heading'], level=1)
            document.add_paragraph(spec['before'])
            table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
            for index, row in enumerate(rows):
                for column, text in enumerate(row):
                    table.cell(index, column).text = text
            document.add_paragraph(spec['after'])
            document.save(stream)
        data = _stable_office(stream.getvalue())
    else:
        raise ValueError('unsupported_fixture_format')
    if len(data) > MAX_SOURCE_BYTES:
        raise ValueError('fixture_source_limit')
    return data
