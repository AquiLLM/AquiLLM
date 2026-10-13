import importlib
from dataclasses import replace
from hashlib import sha256

import pytest
from lib.embeddings.provenance import input_digest
from lib.parsers.source_contracts import SourceContractError
from .test_structural_extraction import extract


class LocalTokenizer:
    revision = 'synthetic-codepoints-v1'
    def __call__(self, text):
        return list(text)


def chunks(extraction, max_tokens=120, overlap_tokens=8, tokenizer=None):
    module = importlib.import_module('lib.parsers.source_chunking')
    tokenizer = tokenizer or LocalTokenizer()
    key = module.chunker_key_for(tokenizer=tokenizer, max_tokens=max_tokens, overlap_tokens=overlap_tokens)
    return module.chunk_extraction(extraction, chunker_key=key, tokenizer=tokenizer, max_tokens=max_tokens, overlap_tokens=overlap_tokens)


def test_contiguous_table_rows_preserve_source_offsets_and_schema():
    result = extract('rows.csv', b'Product,Domestic,Foreign\nWidget,,7\n')
    output = chunks(result)
    assert output
    for chunk in output:
        assert chunk.source_text == result.canonical_text[chunk.start:chunk.end]
        assert len(LocalTokenizer()(chunk.embedding_text)) <= 120
        assert chunk.source_sha256 == sha256(chunk.source_text.encode()).hexdigest()
        assert chunk.input_sha256 == input_digest(chunk.embedding_text)
    assert 'Domestic' in output[-1].embedding_text
    assert all(a.start <= b.start for a, b in zip(output, output[1:]))
    assert set(range(len(result.canonical_text))) == {i for chunk in output for i in range(chunk.start, chunk.end)}


def test_oversized_unicode_element_is_lossless_with_overlap():
    result = extract('notes.txt', ('−4 kg α\n' * 40).encode())
    output = chunks(result, max_tokens=30, overlap_tokens=5)
    assert len(output) > 2
    assert all(len(chunk.embedding_text) <= 30 for chunk in output)
    assert any(a.end > b.start for a, b in zip(output, output[1:]))
    assert set(range(len(result.canonical_text))) == {i for chunk in output for i in range(chunk.start, chunk.end)}


def test_budget_includes_heading_context_and_rejects_unrepresentable_header():
    result = extract('data.csv', b'A very long column name,value\nexample,7\n')
    with pytest.raises(SourceContractError, match='unrepresentable_source_element'):
        chunks(result, max_tokens=5, overlap_tokens=0)


def test_chunker_fingerprint_binds_budget_overlap_and_tokenizer_revision():
    module = importlib.import_module('lib.parsers.source_chunking')
    args = dict(tokenizer=LocalTokenizer(), max_tokens=120, overlap_tokens=8)
    first = module.chunker_key_for(**args)
    assert first != module.chunker_key_for(**dict(args, max_tokens=121))
    assert first != module.chunker_key_for(**dict(args, overlap_tokens=9))
    tokenizer = LocalTokenizer()
    tokenizer.revision = 'synthetic-codepoints-v2'
    assert first != module.chunker_key_for(**dict(args, tokenizer=tokenizer))
    result = extract('notes.txt', b'hello')
    with pytest.raises(SourceContractError, match='chunker_key_mismatch'):
        module.chunk_extraction(result, chunker_key='unknown', **args)


def test_malformed_partition_rejected_before_chunking():
    result = extract('notes.txt', b'hello')
    broken = replace(result, elements=(replace(result.elements[0], start=1),))
    with pytest.raises(SourceContractError, match='invalid_extraction_structure'):
        chunks(broken)


def test_heading_context_formula_qualifiers_and_exact_digest_vectors():
    from docx import Document
    import io
    doc = Document()
    doc.add_heading('Metrics', level=1)
    doc.add_paragraph(''.join(f'−4 kg\nα{number} ' for number in range(20)))
    stream = io.BytesIO()
    doc.save(stream)
    result = extract('metrics.docx', stream.getvalue())
    output = chunks(result, max_tokens=40, overlap_tokens=3)
    assert all(c.embedding_text.startswith('heading=Metrics\n') for c in output[1:])
    assert output[1].source_sha256 != output[2].source_sha256
    assert all(c.source_sha256 != sha256(result.canonical_text.encode()).hexdigest() for c in output)
    # Literal non-ASCII/newline vectors distinguish raw source hashing from
    # established input_digest JSON-string hashing, including escape handling.
    assert sha256('−4 kg\nα'.encode()).hexdigest() == '06dfe8406ad9219197231016bda1adb0913b4e2a4e12b8b55e15e78198d329de'
    assert input_digest('−4 kg\nα') == '1a54d3eb274ea95888db5c2f4c36064fd3b7e4fd0173fafe2986f59564485ae3'
    from .test_structural_extraction import workbook_bytes
    formula = extract('formula.xlsx', workbook_bytes([['Profit', '=50+20']]))
    formula_chunks = chunks(formula)
    assert any('formula==50+20 value=<unknown>' in c.embedding_text for c in formula_chunks)


def test_oversized_formula_row_repeats_unknown_value_qualifiers_in_embedding_context():
    from .test_structural_extraction import workbook_bytes
    result = extract('formulas.xlsx', workbook_bytes([['x' * 150, '=50+20']]))
    output = chunks(result, max_tokens=100, overlap_tokens=4)
    assert len(output) > 1
    assert all('B1 formula==50+20 value=<unknown>' in chunk.embedding_text for chunk in output)
    assert all(chunk.source_text == result.canonical_text[chunk.start:chunk.end] for chunk in output)


def review_bpe_tokenizer(*, suffix=False):
    import tiktoken
    ranks = {bytes([i]): i for i in range(256)}
    ranks.update({b'X\xc3': 256, b'X\xc3\xa9': 257} if suffix else {b'\xa9X': 256, b'\xc3\xa9X': 257})
    name = 'review-byte-bpe-suffix-v1' if suffix else 'review-byte-bpe-v1'
    encoding = tiktoken.Encoding(name=name, pat_str=r'[\s\S]+',
                                 mergeable_ranks=ranks, special_tokens={})
    class Adapter:
        revision = name
        encode = staticmethod(encoding.encode)
    return Adapter()


def test_real_local_bpe_nonmonotone_prefix_remains_representable():
    tokenizer = review_bpe_tokenizer()
    assert [len(tokenizer.encode(text)) for text in ['\u00e9', '\u00e9X', '\u00e9Xq']] == [2, 1, 2]
    result = extract('unicode.txt', '\u00e9Xqrst'.encode())
    output = chunks(result, max_tokens=1, overlap_tokens=0, tokenizer=tokenizer)
    assert ''.join(chunk.source_text for chunk in output) == result.canonical_text
    assert all(len(tokenizer.encode(chunk.embedding_text)) <= 1 for chunk in output)


def test_real_local_bpe_nonmonotone_suffix_preserves_available_overlap():
    tokenizer = review_bpe_tokenizer(suffix=True)
    assert [len(tokenizer.encode(text)) for text in ['\u00e9', 'X\u00e9']] == [2, 1]
    result = extract('unicode.txt', 'abcdX\u00e9qrst'.encode())
    output = chunks(result, max_tokens=5, overlap_tokens=1, tokenizer=tokenizer)
    assert output[0].source_text == 'abcdX\u00e9'
    assert output[1].start == 4
    assert all(len(tokenizer.encode(c.embedding_text)) <= 5 for c in output)
    assert all(c.source_text == result.canonical_text[c.start:c.end] for c in output)
    assert set(range(len(result.canonical_text))) == {i for c in output for i in range(c.start, c.end)}


def test_real_local_bpe_genuinely_unrepresentable_codepoint():
    with pytest.raises(SourceContractError, match='^unrepresentable_source_element$'):
        chunks(extract('unicode.txt', '\u00e9'.encode()), max_tokens=1, overlap_tokens=0,
               tokenizer=review_bpe_tokenizer())


@pytest.mark.parametrize('with_heading', [False, True])
def test_nonmonotone_fallback_has_total_input_codepoint_ceiling(with_heading):
    class NoFit:
        revision = 'adversarial-no-fit-v1'
        lengths = []
        def encode(self, text):
            self.lengths.append(len(text))
            return [0] if text == 'Header\n' else [0, 1]
    tokenizer = NoFit()
    size = 10000
    result = extract('large.txt', (('Header\n' if with_heading else '') + 'x' * size).encode())
    context = 'heading=Header\n' if with_heading else ''
    if with_heading:
        element = result.elements[0]
        result = replace(result, elements=(replace(element, kind='heading', end=7),
                         replace(element, element_key='element-2', start=7)))
    # Pin the ordinary heuristic probes so the fallback work can be measured
    # independently; every tokenizer invocation receives its complete input.
    ordinary = ([7, 7] if with_heading else []) + [size + len(context)]
    low, high = 1, size
    while low <= high:
        middle = (low + high) // 2
        ordinary.append(middle + len(context))
        high = middle - 1
    with pytest.raises(SourceContractError, match='^chunker_work_limit$'):
        chunks(result, max_tokens=1, overlap_tokens=0, tokenizer=tokenizer)
    assert tokenizer.lengths[:len(ordinary)] == ordinary
    fallback = tokenizer.lengths[len(ordinary):]
    assert 0 < len(fallback) <= 4096
    assert sum(fallback) <= 1000000
    assert len(tokenizer.lengths) <= len(ordinary) + 4096


def test_nonmonotone_fallback_call_ceiling_spans_all_elements_without_publication():
    bpe = review_bpe_tokenizer()
    class CountingBPE:
        revision = bpe.revision
        lengths = []
        def encode(self, text):
            self.lengths.append(len(text))
            return bpe.encode(text)
    tokenizer = CountingBPE()
    result = extract('many.txt', ('\u00e9Xqrst' * 1500).encode())
    element = result.elements[0]
    result = replace(result, elements=tuple(replace(element, element_key=f'element-{i+1}',
                     start=i*6, end=(i+1)*6) for i in range(1500)))
    module = importlib.import_module('lib.parsers.source_chunking')
    assert module.MAX_FALLBACK_TOKENIZER_CALLS == 4096
    assert module.MAX_FALLBACK_INPUT_CODEPOINTS == 1000000
    # Each element needs the real BPE fallback; even many already built chunks
    # must never escape the tuple-return boundary when the call ceiling is hit.
    with pytest.raises(SourceContractError, match='^chunker_work_limit$'):
        chunks(result, max_tokens=1, overlap_tokens=0, tokenizer=tokenizer)
    # Each completed element needs exactly three fallback probes (5, 4, 2).
    # The 1366th element gets one probe, reaching 4096, then stops before the
    # next tokenizer call. Five-codepoint inputs occur only in that fallback.
    assert tokenizer.lengths.count(5) == 1366
    assert tokenizer.lengths.count(6) == 1366


def test_chunker_fingerprint_binds_nonmonotone_policy_and_both_work_limits(monkeypatch):
    module = importlib.import_module('lib.parsers.source_chunking')
    options = dict(tokenizer=LocalTokenizer(), max_tokens=120, overlap_tokens=8)
    key = module.chunker_key_for(**options)
    for setting, value in [('FITTING_POLICY', 'different-policy'),
                           ('MAX_FALLBACK_TOKENIZER_CALLS', 4095),
                           ('MAX_FALLBACK_INPUT_CODEPOINTS', 999999)]:
        with monkeypatch.context() as change:
            change.setattr(module, setting, value)
            assert module.chunker_key_for(**options) != key
            with pytest.raises(SourceContractError, match='^chunker_key_mismatch$'):
                module.chunk_extraction(extract('notes.txt', b'hello'), chunker_key=key, **options)
