"""test structural docx extraction: synthetic regression coverage."""

import io
import pytest
from .structural_extraction_support import extract, workbook_bytes, xls_bytes


def test_docx_public_table_fixture_and_nested_order():
    from docx import Document
    doc = Document()
    doc.add_paragraph('Quarterly report')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'SecretRevenue'
    table.cell(0, 1).text = '93017'
    cell = table.cell(0, 1)
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'nested fact'
    cell.add_paragraph('after nested')
    doc.add_paragraph('End report −4 kg')
    output = io.BytesIO()
    doc.save(output)
    result = extract('report.docx', output.getvalue())
    text = result.canonical_text
    assert text.index('Quarterly report') < text.index('SecretRevenue') < text.index('93017') < text.index('nested fact') < text.index('after nested') < text.index('End report −4 kg')
    assert any(e.parent_key for e in result.elements)
    assert ''.join(text[e.start:e.end] for e in result.elements if e.kind != 'figure') == text


def _wrap_docx_control(node):
    from docx.oxml import OxmlElement
    control, content = OxmlElement('w:sdt'), OxmlElement('w:sdtContent')
    node.addprevious(control)
    content.append(node)
    control.append(content)
    return control


def test_docx_body_and_cell_content_controls_preserve_native_order():
    from docx import Document
    doc = Document()
    doc.add_paragraph('Before')
    controlled = doc.add_paragraph('ControlledRevenue 93017')
    _wrap_docx_control(_wrap_docx_control(controlled._p))
    table = doc.add_table(rows=1, cols=2)
    cell = table.cell(0, 0)
    cell.text = 'ControlledCell 481'
    _wrap_docx_control(cell.paragraphs[0]._p)
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'ControlledNested 517'
    _wrap_docx_control(nested._tbl)
    _wrap_docx_control(cell.add_paragraph('ControlledTail 639')._p)
    table.cell(0, 1).text = 'NextCell 752'
    _wrap_docx_control(table._tbl)
    doc.add_paragraph('After')
    output = io.BytesIO()
    doc.save(output)
    result = extract('content-controls.docx', output.getvalue())
    facts = ['Before', 'ControlledRevenue 93017', 'ControlledCell 481',
             'ControlledNested 517', 'ControlledTail 639', 'NextCell 752', 'After']
    positions = [result.canonical_text.index(fact) for fact in facts]
    assert positions == sorted(positions)
    assert result.quality == 'complete'
    assert any(e.parent_key for e in result.elements)
    assert ''.join(result.canonical_text[e.start:e.end] for e in result.elements) == result.canonical_text


@pytest.mark.parametrize('wrapper,inline', [('w:customXml', False), ('w:sdt', True), ('w:sdt', 'row'), ('w:del', True)])
def test_docx_remaining_unsupported_wrappers_are_explicitly_partial(wrapper, inline):
    from docx import Document
    from docx.oxml import OxmlElement
    doc = Document()
    paragraph = doc.add_paragraph('Visible before')
    if inline == 'row':
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'HiddenRowRevenue 93017'
        _wrap_docx_control(table.rows[0]._tr)
    elif inline:
        run = paragraph.add_run('HiddenInlineRevenue 93017')._r
        if wrapper == 'w:sdt':
            _wrap_docx_control(run)
        else:
            from docx.oxml.ns import qn
            run.find(qn('w:t')).tag = qn('w:delText')
            node = OxmlElement(wrapper)
            run.addprevious(node)
            node.append(run)
    else:
        controlled = doc.add_paragraph('HiddenCustomRevenue 93017')
        node = OxmlElement(wrapper)
        controlled._p.addprevious(node)
        node.append(controlled._p)
    doc.add_paragraph('Visible after')
    output = io.BytesIO()
    doc.save(output)
    result = extract('unsupported-wrapper.docx', output.getvalue())
    assert result.quality == 'partial'
    assert 'document_wrapper_structure_unavailable' in result.issue_codes
    assert result.canonical_text.index('Visible before') < result.canonical_text.index('Visible after')


def test_docx_merged_spans_and_native_table_headers():
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    for col, value in enumerate(['Product', 'Domestic', 'Foreign']):
        table.cell(0, col).text = value
    table.cell(1, 0).merge(table.cell(2, 0)).text = 'Widget'
    table.cell(1, 1).text = '7'
    output = io.BytesIO()
    doc.save(output)
    result = extract('table.docx', output.getvalue())
    assert result.elements[1].cells[0].row_span == 2
    assert result.elements[1].cells[1].header_keys == ('Domestic',)
    assert result.elements[2].cells[0].value_status == 'blank'


def test_nested_table_in_first_cell_stays_before_next_cell_text():
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'cell one before'
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'nested first cell'
    table.cell(0, 0).add_paragraph('cell one after')
    table.cell(0, 1).text = 'cell two'
    output = io.BytesIO()
    doc.save(output)
    text = extract('nested.docx', output.getvalue()).canonical_text
    assert text.index('cell one before') < text.index('nested first cell') < text.index('cell one after') < text.index('cell two')


def test_unparsed_docx_auxiliary_content_cannot_claim_complete():
    from docx import Document
    source = Document()
    source.add_paragraph('Body fact')
    source.sections[0].header.paragraphs[0].text = 'Header fact'
    output = io.BytesIO()
    source.save(output)
    result = extract('header.docx', output.getvalue())
    assert result.quality == 'partial'
    assert 'auxiliary_document_content_unavailable' in result.issue_codes
