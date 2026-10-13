"""test structural extraction: synthetic regression coverage."""

import io
import zipfile
import pytest
from lib.parsers.source_contracts import SourceContractError
from .structural_extraction_support import extract, workbook_bytes, xls_bytes
from .test_structural_docx_extraction import _wrap_docx_control


def test_public_wrong_column_fixture_preserves_empty_columns():
    header = ['Product', 'Domestic', 'Foreign']
    a = extract('a.xlsx', workbook_bytes([header, ['Widget', 7, None]]))
    b = extract('b.xlsx', workbook_bytes([header, ['Widget', None, 7]]))
    assert a.canonical_text != b.canonical_text
    assert 'sheet=Sheet row=2 A2 Product=Widget B2 Domestic=7 C2 Foreign=<blank>' in a.canonical_text
    assert a.elements[-1].cells[1].address == 'B2'
    assert b.elements[-1].cells[2].text == '7'


def test_uncached_and_cached_formula_are_distinct():
    data = workbook_bytes([['Profit', '=50+20']])
    uncached = extract('formula.xlsx', data)
    assert 'B1 formula==50+20 value=<unknown>' in uncached.canonical_text
    assert uncached.quality == 'partial'
    assert uncached.elements[-1].cells[1].value_status == 'uncached'
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, 'w') as target:
        for item in source.infolist():
            raw = source.read(item)
            if item.filename == 'xl/worksheets/sheet1.xml':
                raw = raw.replace(b'<v></v>', b'<v>70</v>')
            target.writestr(item, raw)
    cached = extract('formula.xlsx', output.getvalue())
    assert 'B1 formula==50+20 value=70' in cached.canonical_text
    assert cached.elements[-1].cells[1].value_status == 'cached'


def test_csv_quoted_commas_blank_rows_and_trailing_cells():
    result = extract('data.csv', b'Product,Domestic,Foreign\r\n"Widget, large",,7\r\n,,\r\n')
    assert result.elements[1].cells[0].text == 'Widget, large'
    assert result.elements[1].cells[1].value_status == 'blank'
    assert result.elements[1].cells[2].address == 'C2'
    assert len(result.elements[2].cells) == 3


def test_xlsx_chart_only_fact_is_explicitly_partial():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    book = Workbook()
    book.active.append(['Metric', 'Value'])
    book.active.append(['A', 7])
    chart = BarChart()
    chart.title = 'ChartOnlyRevenue 93017'
    chart.add_data(Reference(book.active, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    book.active.add_chart(chart, 'D1')
    output = io.BytesIO()
    book.save(output)
    book.close()
    result = extract('chart.xlsx', output.getvalue())
    assert result.quality == 'partial'
    assert 'spreadsheet_chart_structure_unavailable' in result.issue_codes
    assert result.elements[-1].cells[1].text == '7'


def test_literal_coordinate_markers_cannot_collide_in_canonical_projection():
    a = extract('a.csv', b'"X B1 value=Y",Z\n')
    b = extract('b.csv', b'X,"Y B1 value=Z"\n')
    assert a.canonical_text != b.canonical_text
    assert a.elements[0].cells[0].text == 'X B1 value=Y'
    assert b.elements[0].cells[1].text == 'Y B1 value=Z'


def test_ods_repeated_rows_cells_and_merged_spans():
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableRow, TableCell, CoveredTableCell
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    sheet = Table(name='Metrics')
    row = TableRow(numberrowsrepeated=2)
    cell = TableCell(valuetype='string', numbercolumnsspanned=2)
    cell.addElement(P(text='−4 kg'))
    row.addElement(cell)
    row.addElement(CoveredTableCell())
    row.addElement(TableCell(numbercolumnsrepeated=2))
    sheet.addElement(row)
    doc.spreadsheet.addElement(sheet)
    output = io.BytesIO()
    doc.save(output)
    result = extract('metrics.ods', output.getvalue())
    assert [e.cells[0].address for e in result.elements] == ['A1', 'A2']
    assert result.elements[0].cells[0].column_span == 2
    assert [c.address for c in result.elements[0].cells] == ['A1', 'B1', 'C1', 'D1']
    assert result.elements[0].cells[1].value_status == 'blank'


def test_xls_blank_columns_and_merged_spans():
    result = extract('legacy.xls', xls_bytes())
    assert result.elements[-1].cells[1].value_status == 'blank'
    assert result.elements[-1].cells[2].text == '7'
    assert result.elements[0].cells[0].column_span == 2
    assert result.quality == 'partial'  # xlrd cannot prove formula/cache identity.


def test_pptx_native_table_grouped_text_and_order():
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'Before'
    table = slide.shapes.add_table(1, 2, 0, Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = 'NativeFact'
    table.cell(0, 1).text = '−4 kg'
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, Inches(2), Inches(2), Inches(1)).text = 'Grouped fact'
    output = io.BytesIO()
    deck.save(output)
    result = extract('deck.pptx', output.getvalue())
    assert result.canonical_text.index('Before') < result.canonical_text.index('NativeFact') < result.canonical_text.index('Grouped fact')
    assert all(e.slide == 1 for e in result.elements)
    assert any(e.parent_key for e in result.elements)


@pytest.mark.parametrize('extension', ['doc', 'ppt', 'rtf', 'odt'])
def test_legacy_and_unparsed_formats_never_decode_garbage_as_complete(extension):
    result = extract('legacy.' + extension, b'garbage binary\xff')
    assert result.quality == 'unsupported'
    assert result.canonical_text == ''


def test_pdf_page_identity_and_unknown_layout():
    import fitz
    doc = fitz.open()
    for text in ['First page', 'Second page']:
        doc.new_page().insert_text((30, 30), text)
    result = extract('pages.pdf', doc.tobytes())
    doc.close()
    assert [e.page for e in result.elements] == [1, 2]
    assert all(e.bbox is None for e in result.elements)
    assert result.quality == 'partial'
    assert 'layout_unavailable' in result.issue_codes


def test_malformed_source_has_stable_issue_code():
    with pytest.raises(SourceContractError, match='malformed_source_structure'):
        extract('broken.docx', b'broken')


def test_xlsx_both_readers_close_on_cache_reader_failure(monkeypatch):
    from openpyxl import load_workbook
    import openpyxl
    data = workbook_bytes([['value', '=50+20']])
    opened = []
    original = load_workbook
    def load(*args, **kwargs):
        if kwargs['data_only']:
            raise ValueError('synthetic cache-reader failure')
        book = original(*args, **kwargs)
        opened.append(book)
        close = book.close
        book.closed_by_extractor = False
        def tracking_close():
            book.closed_by_extractor = True
            close()
        book.close = tracking_close
        return book
    monkeypatch.setattr(openpyxl, 'load_workbook', load)
    with pytest.raises(SourceContractError, match='malformed_source_structure'):
        extract('workbook.xlsx', data)
    assert opened[0].closed_by_extractor


def test_xlsx_merge_and_blank_cells_have_explicit_coordinates():
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['Product', 'Domestic', 'Foreign'])
    book.active.append(['Widget', None, 7])
    book.active.merge_cells('A2:A3')
    output = io.BytesIO()
    book.save(output)
    book.close()
    result = extract('merged.xlsx', output.getvalue())
    assert result.elements[1].cells[0].row_span == 2
    assert result.elements[2].cells[0].address == 'A3'
    assert result.elements[2].cells[0].value_status == 'blank'


def test_blank_pdf_retains_page_location_without_inventing_body_text():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    result = extract('blank.pdf', output.getvalue())
    assert result.canonical_text == ''
    assert result.elements[0].page == 1
    assert result.elements[0].start == result.elements[0].end == 0


@pytest.mark.parametrize('format', ['docx', 'pptx', 'xlsx', 'ods'])
def test_embedded_images_have_figure_overlays_with_native_locations(format):
    from apps.ingestion.tests.test_office_conversion_fallback import image_bytes
    png = image_bytes()
    output = io.BytesIO()
    if format == 'docx':
        from docx import Document
        source = Document()
        source.add_paragraph('Native body')
        source.add_picture(io.BytesIO(png))
        source.save(output)
    elif format == 'pptx':
        from pptx import Presentation
        source = Presentation()
        slide = source.slides.add_slide(source.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = 'Native body'
        slide.shapes.add_picture(io.BytesIO(png), 0, 0)
        source.save(output)
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image
        source = Workbook()
        source.active.append(['Native body'])
        source.active.add_image(Image(io.BytesIO(png)), 'B2')
        source.save(output)
        source.close()
    else:
        from odf.opendocument import OpenDocumentSpreadsheet
        from odf.table import Table, TableRow, TableCell
        from odf.draw import Frame, Image
        from odf.text import P
        source = OpenDocumentSpreadsheet()
        table = Table(name='Sheet')
        row = TableRow()
        cell = TableCell()
        cell.addElement(P(text='Native body'))
        frame = Frame(width='1cm', height='1cm')
        frame.addElement(Image(href='Pictures/fixture.png', type='simple'))
        cell.addElement(frame)
        row.addElement(cell)
        table.addElement(row)
        source.spreadsheet.addElement(table)
        source.addPicture('Pictures/fixture.png', mediatype='image/png', content=png)
        source.save(output)
    result = extract('image.' + format, output.getvalue())
    figures = [e for e in result.elements if e.kind == 'figure']
    assert len(figures) == 1
    assert figures[0].start == figures[0].end
    assert ''.join(result.canonical_text[e.start:e.end] for e in result.elements if e.kind != 'figure') == result.canonical_text
    if format == 'pptx':
        assert figures[0].slide == 1
        assert figures[0].bbox is not None
    if format in ('xlsx', 'ods'):
        assert figures[0].sheet == 'Sheet'


def test_unparsed_pptx_charts_cannot_claim_complete():
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    data = ChartData()
    data.categories = ['Metric']
    data.add_series('value', (7,))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, 1000000, 1000000, data)
    output = io.BytesIO()
    source.save(output)
    result = extract('chart.pptx', output.getvalue())
    assert result.quality == 'partial'
    assert 'slide_shape_structure_unavailable' in result.issue_codes
