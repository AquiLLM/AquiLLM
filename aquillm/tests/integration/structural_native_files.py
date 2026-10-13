"""Public native byte fixtures; application execution is deliberately still pending."""
from io import BytesIO

BODY = ('Measurements in 2028, 2029, and 2030 followed Morrow (2027) and Chen (2026). '
        'The signed concentration was 7 ug/L only at pH 6; measurements above pH 6 were excluded.')
REFERENCE = 'Morrow, A. (2031). Synthetic comet calibration. Journal of Example Measurements, 4, 17.'
HEADER = 'Concentration (ug/L); signed result only'
SIGNED = 'CERULEAN-OTTER records 7 ug/L only at pH 6.'
OPPOSITION = 'The draft 9 ug/L is contradicted; measurements above pH 6 were excluded.'
LONG = ('The draft prefix is unsigned and reports no final measurement. ' * 28
        + SIGNED + ' ' + OPPOSITION)


def document_bytes(*, late=False):
    from docx import Document
    document = Document()
    if late:
        document.add_heading(HEADER, level=1)
        document.add_paragraph(LONG)
    else:
        document.add_heading('Results', level=1)
        document.add_paragraph(BODY)
        document.add_heading('References', level=1)
        document.add_paragraph(REFERENCE)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def scrambled_pdf_bytes():
    """Actual two-column positions, deliberately interleaved content-stream order."""
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
        DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    operations = []
    for x, y, text in (
        (36, 740, 'References'), (320, 740, 'Results'),
        (36, 720, 'Morrow (2031). Synthetic comet calibration.'),
        (320, 720, 'Measurements in 2028, 2029, 2030 followed Morrow (2027).'),
        (36, 700, 'Journal of Example Measurements, 4, 17.'),
        (320, 700, 'The signed concentration was 7 ug/L only at pH 6;'),
        (320, 680, 'measurements above pH 6 were excluded.'),
    ):
        literal = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
        operations.append(f'BT /F1 8 Tf {x} {y} Td ({literal}) Tj ET')
    stream = DecodedStreamObject()
    stream.set_data('\n'.join(operations).encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    result = BytesIO()
    writer.write(result)
    writer.close()
    return result.getvalue()
