"""Offline adversarial ingestion probes; no Django setup, database, or HTTP.

Run from repository root with:
    rtk proxy python -B docs/audits/2026-10-10-rag-adversarial/ingestion_probes.py

Uses installed python-docx/openpyxl for in-memory fixtures. Ingestion control
flow runs from the current AST against fake persistence, not production models.
"""

from __future__ import annotations

import ast
import io
from pathlib import Path
import runpy
from types import SimpleNamespace as NS


ROOT = Path(__file__).resolve().parents[3]


def load_function(path: str, name: str, namespace: dict):
    tree = ast.parse((ROOT / path).read_text(encoding="utf-8"))
    function = next(
        node for node in tree.body
        if isinstance(node, ast.FunctionDef) and node.name == name
    )
    function.decorator_list = []
    function.body = [node for node in function.body if not isinstance(node, (ast.Import, ast.ImportFrom))]
    exec(compile(ast.Module(body=[function], type_ignores=[]), path, "exec"), namespace)
    return namespace[name]


def parser_probes():
    from docx import Document
    from openpyxl import Workbook

    docx_extract = runpy.run_path(str(ROOT / "aquillm/lib/parsers/documents/docx.py"))["extract_docx_text"]
    xlsx_extract = runpy.run_path(str(ROOT / "aquillm/lib/parsers/spreadsheets/xlsx.py"))["extract_xlsx_text"]
    document = Document()
    document.add_paragraph("Quarterly report")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "SecretRevenue"
    table.cell(0, 1).text = "93017"
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = docx_extract(buffer.getvalue())
    assert extracted == "Quarterly report", extracted
    print("CONFIRMED: DOCX native table SecretRevenue=93017 disappears:", repr(extracted))

    def workbook_text(rows):
        workbook = Workbook()
        for row in rows:
            workbook.active.append(row)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return xlsx_extract(buffer.getvalue())

    header = ["Product", "Domestic", "Foreign"]
    domestic = workbook_text([header, ["Widget", 7, None]])
    foreign = workbook_text([header, ["Widget", None, 7]])
    assert domestic == foreign, (domestic, foreign)
    print("CONFIRMED: XLSX Domestic=7 and Foreign=7 serialize identically:", repr(domestic))
    formula = workbook_text([["Profit", "=50+20"]])
    assert formula == "# Sheet: Sheet\nProfit", formula
    print("CONFIRMED: Uncached XLSX formula disappears:", repr(formula))


def ingestion_probes():
    saved_rows = []
    saved_files = []

    class FakeIntegrityError(Exception):
        pass

    class FakeContentFile(io.BytesIO):
        pass

    class FakeDocument:
        def __init__(self, **kwargs):
            self.__dict__.update(kwargs)
            self.id = len(saved_rows) + 1
            self.pdf_file = NS(save=lambda filename, content, save=False: saved_files.append((filename, content.read())))

        def save(self):
            if self.title == "duplicate":
                raise FakeIntegrityError("existing collection/text hash collision")
            saved_rows.append(self.title)

    manager = NS(
        select_related=lambda *args: manager,
        filter=lambda **kwargs: manager,
        first=lambda: item,
    )
    namespace = runpy.run_path(str(ROOT / "aquillm/aquillm/task_ingest_helpers.py"))
    namespace.update(
        IngestionBatchItem=NS(objects=manager, Status=NS(PROCESSING="processing", SUCCESS="success", ERROR="error")),
        timezone=NS(now=lambda: "offline-time"),
        ContentFile=FakeContentFile,
        IntegrityError=FakeIntegrityError,
        Collection=NS(),
        CollectionPermission=NS(),
        DocumentFigure=type("FakeFigure", (), {}),
        ImageUploadDocument=type("FakeImage", (), {}),
        MediaUploadDocument=type("FakeMedia", (), {}),
        PDFDocument=FakeDocument,
        RawTextDocument=FakeDocument,
    )
    payload = lambda title, kind: NS(
        title=title, normalized_type=kind, full_text="extracted PDF text",
        modality="text", metadata={}, provider="", model="", media_content_type="",
    )
    item = NS(
        id=1, source_file=NS(open=lambda mode: io.BytesIO(b"PK outer ZIP bytes")),
        original_filename="bundle.zip", content_type="application/zip",
        batch=NS(collection=NS(), user=NS()), save=lambda **kwargs: None,
        document_ids=[],
    )
    namespace["extract_text_payloads"] = lambda *args, **kwargs: [payload("paper", "pdf")]
    ingest = load_function("aquillm/aquillm/task_ingest_uploaded.py", "run_ingest_uploaded_file", namespace)
    ingest(1)
    assert saved_files == [("bundle.zip", b"PK outer ZIP bytes")], saved_files
    assert item.status == "success", item.status
    print("CONFIRMED: PDF payload from ZIP persists outer ZIP bytes/name and reports success:", saved_files)

    saved_rows.clear()
    item.document_ids = []
    namespace["extract_text_payloads"] = lambda *args, **kwargs: [payload("first", "txt"), payload("duplicate", "txt")]
    ingest(1)
    assert saved_rows == ["first"] and item.status == "error" and item.document_ids == []
    print("CONFIRMED: Later collision leaves earlier document saved, item error, document_ids empty:", saved_rows, item.status, item.document_ids)


def embedding_reuse_probe():
    """Exercise current duplicate lookup with valid old-space receipt mocked in.

    The receipt validation seam is mocked; this probe concerns whether the
    duplicate-selection policy asks about desired embedding/chunker identity.
    """
    filters = []
    old_receipt = {"declared_identity": {"model": "old-model"}}
    old_chunk = NS(
        content="same text", start_position=0, end_position=9, chunk_number=0,
        modality="text", metadata={}, embedding=[1.0, 0.0], embedding_provenance=old_receipt,
    )
    duplicate = NS(id="old-document")
    model_manager = NS()
    model_manager.using = lambda alias: model_manager
    model_manager.filter = lambda **kwargs: (filters.append(kwargs) or model_manager)
    model_manager.exclude = lambda **kwargs: model_manager
    model_manager.order_by = lambda *args: model_manager
    model_manager.first = lambda: duplicate
    chunk_manager = NS()
    chunk_manager.using = lambda alias: chunk_manager
    chunk_manager.filter = lambda **kwargs: chunk_manager
    chunk_manager.order_by = lambda *args: [old_chunk]
    class FakeChunk(NS):
        objects = chunk_manager
        Modality = NS(TEXT="text")
    namespace = dict(
        TextChunk=FakeChunk,
        DESCENDED_FROM_DOCUMENT=[NS(_base_manager=model_manager)],
        valid_provenance=lambda vector, receipt: receipt,
    )
    duplicate_chunks = load_function("aquillm/apps/documents/tasks/chunking.py", "_duplicate_chunks", namespace)
    result = duplicate_chunks(NS(id="new-document", full_text_hash="same-content-hash"), using="offline")
    assert filters == [{"full_text_hash": "same-content-hash", "ingestion_complete": True}], filters
    assert result[0].embedding_provenance["declared_identity"]["model"] == "old-model"
    print("CONFIRMED: Duplicate reuse asks only text hash/completion and carries old-model receipt:", filters)


if __name__ == "__main__":
    parser_probes()
    ingestion_probes()
    embedding_reuse_probe()
    print("All offline probes reproduced the reviewed behavior; no database or provider calls.")
